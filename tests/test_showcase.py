"""Tests for the portfolio showcase runner and demonstration summary generator."""

import hashlib
import json
from pathlib import Path
import openpyxl
import pytest

from retail_analytics.showcase.summary import get_project_root, run_showcase_demo


def _compute_file_sha256(path: Path) -> str:
    """Compute the SHA-256 hash of a file."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest()


def test_showcase_run_success():
    """Verify that parameterless run_showcase_demo executes on the bundled synthetic fixture."""
    summary = run_showcase_demo()

    assert summary["showcase_name"] == "UMKM Retail Analytics Technical Showcase"
    assert summary["provenance"]["source"] == "Designated Bundled Synthetic Fixture"
    assert summary["data_quality_audit"]["raw_rows"] == 1788
    assert summary["data_quality_audit"]["accepted_rows"] == 1770
    assert summary["data_quality_audit"]["quarantined_rows"] == 12
    assert summary["data_quality_audit"]["duplicate_rows"] == 6
    assert summary["sales_kpis"]["net_sales"] == 50165500.0
    assert summary["sales_kpis"]["unique_orders"] == 954

    # Verify lightweight JSON artifact was created in fixed output directory
    repo_root = get_project_root()
    summary_file = repo_root / "output" / "showcase" / "showcase_summary.json"
    assert summary_file.exists()

    with open(summary_file, "r", encoding="utf-8") as f:
        loaded = json.load(f)
    assert loaded["sales_kpis"]["net_sales"] == summary["sales_kpis"]["net_sales"]


def test_showcase_reproducibility():
    """Verify deterministic calculations across consecutive parameterless runs."""
    res1 = run_showcase_demo()
    res2 = run_showcase_demo()

    assert res1["sales_kpis"]["net_sales"] == res2["sales_kpis"]["net_sales"]
    assert res1["sales_kpis"]["unique_orders"] == res2["sales_kpis"]["unique_orders"]
    assert res1["data_quality_audit"]["accepted_rows"] == res2["data_quality_audit"]["accepted_rows"]
    assert res1["data_quality_audit"]["quarantined_rows"] == res2["data_quality_audit"]["quarantined_rows"]


def test_showcase_never_overwrites_sample_report_sha256():
    """Regression test proving by SHA-256 hash that the committed report is never modified or overwritten."""
    repo_root = get_project_root()
    sample_report = repo_root / "examples" / "sample_report.xlsx"
    assert sample_report.exists()

    hash_before = _compute_file_sha256(sample_report)
    mtime_before = sample_report.stat().st_mtime
    size_before = sample_report.stat().st_size
    files_before = set(sample_report.parent.iterdir())

    # Execute showcase demo
    run_showcase_demo()

    hash_after = _compute_file_sha256(sample_report)
    mtime_after = sample_report.stat().st_mtime
    size_after = sample_report.stat().st_size
    files_after = set(sample_report.parent.iterdir())

    assert hash_before == hash_after
    assert mtime_before == mtime_after
    assert size_before == size_after
    assert files_before == files_after


def test_sample_report_xlsx_openable_and_valid():
    """Verify that the committed example report exists, is valid Excel, and contains all 8 sheets."""
    repo_root = get_project_root()
    sample_report = repo_root / "examples" / "sample_report.xlsx"
    assert sample_report.exists()

    wb = openpyxl.load_workbook(sample_report, data_only=True)
    expected_sheets = [
        "01_EXECUTIVE_SUMMARY",
        "02_SALES_OVERVIEW",
        "03_PRODUCT_ANALYSIS",
        "04_CUSTOMER_RFM",
        "05_MARKET_BASKET",
        "06_INVENTORY",
        "07_DATA_QUALITY",
        "08_RAW_DATA",
    ]
    assert wb.sheetnames == expected_sheets
    ws = wb["01_EXECUTIVE_SUMMARY"]
    subtitle = ws.cell(row=2, column=1).value
    assert "Synthetic Demonstration Dataset" in subtitle
